import io
import csv
from datetime import datetime
from typing import Dict, Any, List, Optional
from sqlalchemy.orm import Session

from app.models import (
    Vehicle, Driver, Shipment, Trip, MaintenanceLog, FuelLog,
    VehicleStatus, ShipmentStatus, TripStatus
)
from app.services.analytics_service import AnalyticsService

# ReportLab imports for professional PDF generation
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# openpyxl import for Excel generation
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ReportsService:
    @staticmethod
    def get_fleet_utilization_data(db: Session) -> Dict[str, Any]:
        vehicles = db.query(Vehicle).all()
        trips = db.query(Trip).all()
        shipments = db.query(Shipment).all()
        
        utilization = AnalyticsService.get_fleet_utilization(db)
        overview = AnalyticsService.get_operational_overview(db)
        
        rows = []
        for v in vehicles:
            rows.append({
                "vehicle_id": v.vehicle_id,
                "registration": v.registration_number,
                "type": v.vehicle_type,
                "capacity_tons": v.capacity,
                "fuel_type": v.fuel_type,
                "status": v.status.value if hasattr(v.status, "value") else str(v.status),
                "odometer_km": v.odometer_km or 0.0
            })
            
        return {
            "title": "Fleet Utilization & Asset Availability Report",
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
            "summary": {
                "total_fleet_size": overview["total_fleet_size"],
                "active_fleet_count": overview["active_fleet_count"],
                "maintenance_fleet_count": overview["maintenance_fleet_count"],
                "fleet_utilization_rate": f"{overview['fleet_utilization_rate']}%"
            },
            "columns": ["Vehicle ID", "Registration", "Type", "Capacity (T)", "Fuel Type", "Status", "Odometer (km)"],
            "data": rows
        }

    @staticmethod
    def get_fuel_consumption_data(db: Session) -> Dict[str, Any]:
        fuel_analytics = AnalyticsService.get_fuel_analytics(db)
        fuel_logs = db.query(FuelLog).order_by(FuelLog.logged_at.desc()).limit(100).all()
        
        rows = []
        for log in fuel_logs:
            rows.append({
                "vehicle_id": log.vehicle_id,
                "liters": log.liters_filled,
                "cost_per_liter": f"${log.cost_per_liter:.2f}",
                "total_cost": f"${log.total_cost:.2f}",
                "odometer": log.odometer_reading,
                "station": log.fuel_station or "N/A",
                "anomaly": "ANOMALY" if log.is_anomaly else "NORMAL",
                "date": log.logged_at.strftime("%Y-%m-%d %H:%M") if log.logged_at else "N/A"
            })
            
        return {
            "title": "Fleet Fuel Consumption & Efficiency Report",
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
            "summary": {
                "total_fuel_consumed": f"{fuel_analytics['total_liters']} L",
                "total_spend": f"${fuel_analytics['total_fuel_cost']:.2f}",
                "average_efficiency": f"{fuel_analytics['average_efficiency_km_per_l']} km/L",
                "anomalies_flagged": fuel_analytics["anomalies_detected"]
            },
            "columns": ["Vehicle ID", "Liters", "Price/L", "Total Cost", "Odometer", "Station", "Status", "Date"],
            "data": rows
        }

    @staticmethod
    def get_driver_performance_data(db: Session) -> Dict[str, Any]:
        drivers = db.query(Driver).order_by(Driver.safety_score.desc()).all()
        rows = []
        for d in drivers:
            rows.append({
                "driver_code": d.driver_code,
                "name": d.name,
                "license_type": d.license_type,
                "status": d.status.value if hasattr(d.status, "value") else str(d.status),
                "rating": f"{d.rating:.1f}/5.0",
                "safety_score": f"{d.safety_score:.1f}%",
                "trips": d.total_trips,
                "vehicle": d.current_vehicle_id or "Unassigned"
            })
            
        return {
            "title": "Commercial Driver Performance & Safety Audit Report",
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
            "summary": {
                "total_registered_drivers": len(drivers),
                "active_drivers_on_duty": sum(1 for d in drivers if (hasattr(d.status, "value") and d.status.value in ["On Duty", "On Trip"])),
                "average_safety_index": f"{(sum(d.safety_score for d in drivers) / len(drivers)):.1f}%" if drivers else "0.0%"
            },
            "columns": ["Driver ID", "Name", "License", "Duty Status", "Rating", "Safety Score", "Trips", "Assigned Asset"],
            "data": rows
        }

    @staticmethod
    def get_delivery_performance_data(db: Session) -> Dict[str, Any]:
        shipments = db.query(Shipment).order_by(Shipment.created_at.desc()).limit(100).all()
        overview = AnalyticsService.get_operational_overview(db)
        
        rows = []
        for s in shipments:
            rows.append({
                "tracking_number": s.tracking_number,
                "origin": s.origin,
                "destination": s.destination,
                "weight_kg": s.weight_kg,
                "status": s.status.value if hasattr(s.status, "value") else str(s.status),
                "vehicle": s.vehicle_id or "Unassigned",
                "eta": s.eta or "Calculated Live",
                "distance_km": f"{s.distance_km:.1f} km" if s.distance_km else "N/A"
            })
            
        return {
            "title": "Logistics Consignment & Delivery Performance Report",
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
            "summary": {
                "on_time_delivery_rate": f"{overview['on_time_delivery_rate']}%",
                "total_shipments_delivered": overview["total_shipments_delivered"],
                "total_distance_covered": f"{overview['total_distance_km']} km"
            },
            "columns": ["Tracking #", "Origin", "Destination", "Weight (kg)", "Status", "Vehicle", "ETA", "Distance"],
            "data": rows
        }

    @staticmethod
    def get_maintenance_data(db: Session) -> Dict[str, Any]:
        logs = db.query(MaintenanceLog).order_by(MaintenanceLog.scheduled_date.desc()).all()
        summary = AnalyticsService.get_operational_overview(db)
        
        rows = []
        for m in logs:
            rows.append({
                "job_id": m.job_id,
                "vehicle_id": m.vehicle_id,
                "category": m.category,
                "priority": m.priority.value if hasattr(m.priority, "value") else str(m.priority),
                "status": m.status.value if hasattr(m.status, "value") else str(m.status),
                "cost": f"${(m.actual_cost or m.estimated_cost):.2f}",
                "scheduled": m.scheduled_date.strftime("%Y-%m-%d") if m.scheduled_date else "N/A",
                "facility": m.service_center
            })
            
        return {
            "title": "Fleet Preventive & Corrective Maintenance Report",
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
            "summary": {
                "total_service_jobs": len(logs),
                "total_maintenance_spend": f"${summary['total_maintenance_spend']:.2f}",
                "active_alerts": summary["active_maintenance_alerts"]
            },
            "columns": ["Job ID", "Vehicle ID", "Category", "Priority", "Status", "Cost", "Date", "Depot"],
            "data": rows
        }

    @classmethod
    def get_report_data(cls, report_type: str, db: Session) -> Dict[str, Any]:
        rtype = report_type.lower().replace("-", "_")
        if "utilization" in rtype or "fleet" in rtype:
            return cls.get_fleet_utilization_data(db)
        elif "fuel" in rtype:
            return cls.get_fuel_consumption_data(db)
        elif "driver" in rtype:
            return cls.get_driver_performance_data(db)
        elif "delivery" in rtype or "shipment" in rtype:
            return cls.get_delivery_performance_data(db)
        elif "maint" in rtype:
            return cls.get_maintenance_data(db)
        else:
            return cls.get_fleet_utilization_data(db)

    @classmethod
    def generate_csv(cls, report_data: Dict[str, Any]) -> str:
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Metadata Header
        writer.writerow(["FLEETFLOW ENTERPRISE LOGISTICS REPORT"])
        writer.writerow(["Report Title", report_data["title"]])
        writer.writerow(["Generated At", report_data["generated_at"]])
        writer.writerow([])
        
        # Summary Section
        writer.writerow(["EXECUTIVE SUMMARY"])
        for k, v in report_data.get("summary", {}).items():
            writer.writerow([k.replace("_", " ").title(), str(v)])
        writer.writerow([])
        
        # Table Data
        writer.writerow(report_data["columns"])
        for row in report_data["data"]:
            writer.writerow(list(row.values()))
            
        return output.getvalue()

    @classmethod
    def generate_excel(cls, report_data: Dict[str, Any]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"
        
        # Styles
        title_font = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
        section_font = Font(name="Segoe UI", size=12, bold=True, color="1F2937")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="Segoe UI", size=10)
        
        header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        summary_fill = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Title
        ws.cell(row=1, column=1, value="FLEETFLOW ENTERPRISE LOGISTICS REPORT").font = title_font
        ws.cell(row=2, column=1, value=report_data["title"]).font = Font(name="Segoe UI", size=13, bold=True, color="2563EB")
        ws.cell(row=3, column=1, value=f"Generated At: {report_data['generated_at']}").font = sub_font
        
        # Summary
        ws.cell(row=5, column=1, value="EXECUTIVE METRICS").font = section_font
        cur_row = 6
        for k, v in report_data.get("summary", {}).items():
            c1 = ws.cell(row=cur_row, column=1, value=k.replace("_", " ").title())
            c2 = ws.cell(row=cur_row, column=2, value=str(v))
            c1.font = Font(name="Segoe UI", size=10, bold=True)
            c2.font = cell_font
            c1.fill = summary_fill
            c2.fill = summary_fill
            c1.border = thin_border
            c2.border = thin_border
            cur_row += 1
            
        cur_row += 2
        # Data Header
        cols = report_data["columns"]
        for col_idx, col_name in enumerate(cols, 1):
            cell = ws.cell(row=cur_row, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            
        cur_row += 1
        # Data Rows
        for r_idx, row_dict in enumerate(report_data["data"]):
            for c_idx, val in enumerate(row_dict.values(), 1):
                cell = ws.cell(row=cur_row, column=c_idx, value=val)
                cell.font = cell_font
                cell.border = thin_border
                if r_idx % 2 == 1:
                    cell.fill = alt_fill
            cur_row += 1
            
        # Adjust Column Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @classmethod
    def generate_pdf(cls, report_data: Dict[str, Any]) -> bytes:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Normal'],
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#475569'),
            spaceAfter=12
        )
        section_style = ParagraphStyle(
            'SectionHeader',
            parent=styles['Heading2'],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#0f172a'),
            spaceBefore=8,
            spaceAfter=6
        )
        cell_style = ParagraphStyle(
            'TableCell',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#1e293b')
        )
        header_cell_style = ParagraphStyle(
            'TableHeaderCell',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            fontName="Helvetica-Bold",
            textColor=colors.white
        )

        elements = []
        
        # Header Banner
        elements.append(Paragraph("FleetFlow Logistics Operations", title_style))
        elements.append(Paragraph(f"<b>{report_data['title']}</b> | Generated: {report_data['generated_at']}", subtitle_style))
        elements.append(Spacer(1, 8))
        
        # Executive Summary Section
        elements.append(Paragraph("Executive Performance Summary", section_style))
        summary_rows = [["Metric", "Value"]]
        for k, v in report_data.get("summary", {}).items():
            summary_rows.append([k.replace("_", " ").title(), str(v)])
            
        summary_table = Table(summary_rows, colWidths=[200, 340])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563eb')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1'))
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 14))
        
        # Detailed Records Table
        elements.append(Paragraph("Operational Log Details", section_style))
        cols = report_data["columns"]
        header_cells = [Paragraph(c, header_cell_style) for c in cols]
        table_data = [header_cells]
        
        for row in report_data.get("data", [])[:40]:  # limit to 40 records on PDF to ensure clean fit
            row_cells = [Paragraph(str(v), cell_style) for v in row.values()]
            table_data.append(row_cells)
            
        col_width = 540 / len(cols)
        records_table = Table(table_data, colWidths=[col_width] * len(cols))
        records_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0'))
        ]))
        elements.append(records_table)
        
        doc.build(elements)
        return buffer.getvalue()
